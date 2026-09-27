"""Unit tests for document parsers and text cleaning."""
from io import BytesIO

import pytest
from docx import Document as DocxDocument

from app.utils import parsers
from app.utils.parsers import clean_text, parse_docx, parse_file, parse_pdf, parse_txt


def test_clean_text_normalizes_whitespace():
    raw = "First line\r\nSecond   line\r\n\r\n\r\n\r\nThird\tline\x00"
    # `_drop_garbage_lines` collapses blank lines, so paragraphs join here;
    # the test still verifies CRLF/CR normalization, space/tab collapse and
    # control-character removal.
    assert clean_text(raw) == "First line\nSecond line\nThird line"


def test_parse_txt_decodes_and_cleans():
    parsed = parse_txt(b"Hello\r\n\r\n\r\nWorld")
    assert parsed.pages[0].page_number == 1
    assert parsed.text == "Hello\nWorld"


def test_parse_docx_extracts_paragraphs_and_tables():
    doc = DocxDocument()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell A"
    table.rows[0].cells[1].text = "Cell B"

    buffer = BytesIO()
    doc.save(buffer)

    parsed = parse_docx(buffer.getvalue())
    assert "First paragraph" in parsed.text
    assert "Second paragraph" in parsed.text
    assert "Cell A" in parsed.text
    assert "Cell B" in parsed.text
    assert parsed.pages[0].page_number == 1


def test_parse_pdf_extracts_and_numbers_pages(monkeypatch):
    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        def __init__(self, stream) -> None:
            self.pages = [FakePage("First page\r\ncontent"), FakePage("Second page\n\ncontent")]

    monkeypatch.setattr(parsers, "PdfReader", FakeReader)

    parsed = parse_pdf(b"fake-bytes")
    assert len(parsed.pages) == 2
    assert parsed.pages[0].page_number == 1
    assert "First page" in parsed.pages[0].text
    assert "Second page" in parsed.pages[1].text


def test_parse_pdf_skips_empty_pages(monkeypatch):
    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        def __init__(self, stream) -> None:
            self.pages = [FakePage(""), FakePage("real content")]

    monkeypatch.setattr(parsers, "PdfReader", FakeReader)

    parsed = parse_pdf(b"fake-bytes")
    assert len(parsed.pages) == 1
    assert parsed.pages[0].page_number == 2


def test_parse_file_dispatches_by_extension():
    parsed = parse_file("note.txt", b"plain text")
    assert parsed.text == "plain text"


def test_parse_file_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_file("archive.zip", b"data")
